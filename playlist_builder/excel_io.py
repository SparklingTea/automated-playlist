"""Saves matched tracks to an .xlsx file and loads them back, so a list only
has to be matched against Spotify once."""
import io
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from playlist_builder.parser import SongQuery

HEADERS = ["Input", "Title", "Artist", "Spotify URI", "Spotify link"]
_TRACK_ID_RE = re.compile(r"(?:spotify:track:|open\.spotify\.com/(?:intl-[\w-]+/)?track/)([A-Za-z0-9]{22})")


class ExcelFormatError(ValueError):
    pass


def to_excel_bytes(matches: list[tuple[SongQuery, dict | None]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Playlist"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for query, track in matches:
        if track:
            track_id = track["uri"].rsplit(":", 1)[-1]
            ws.append([
                query.raw_line,
                track["name"],
                ", ".join(a["name"] for a in track["artists"]),
                track["uri"],
                f"https://open.spotify.com/track/{track_id}",
            ])
        else:
            ws.append([query.raw_line, None, None, None, None])

    for column, width in zip("ABCDE", (45, 35, 30, 40, 55)):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def from_excel(file) -> list[tuple[SongQuery, dict | None]]:
    """Reads rows back into (query, track) pairs. Rows without a valid
    Spotify URI or link come back unmatched."""
    ws = load_workbook(file, read_only=True, data_only=True).active
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip().lower() if h is not None else "" for h in next(rows, [])]

    def col(name: str) -> int | None:
        return header.index(name.lower()) if name.lower() in header else None

    uri_col, link_col = col("Spotify URI"), col("Spotify link")
    if uri_col is None and link_col is None:
        raise ExcelFormatError("The sheet needs a 'Spotify URI' or 'Spotify link' column in its first row.")
    input_col, title_col, artist_col = col("Input"), col("Title"), col("Artist")

    def cell(row, index):
        if index is None or index >= len(row) or row[index] is None:
            return ""
        return str(row[index]).strip()

    matches = []
    for row in rows:
        title, artist = cell(row, title_col), cell(row, artist_col)
        raw_line = cell(row, input_col) or " - ".join(p for p in (title, artist) if p)
        id_match = _TRACK_ID_RE.search(cell(row, uri_col)) or _TRACK_ID_RE.search(cell(row, link_col))
        if not raw_line and not id_match:
            continue

        track = None
        if id_match:
            track = {
                "name": title or raw_line,
                "artists": [{"name": a.strip()} for a in artist.split(",") if a.strip()],
                "uri": f"spotify:track:{id_match.group(1)}",
            }
        matches.append((SongQuery(raw_line=raw_line or track["uri"], title=title or None, artist=artist or None), track))
    return matches
